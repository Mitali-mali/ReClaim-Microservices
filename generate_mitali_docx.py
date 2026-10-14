from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_mitali_docx():
    docx_path = r"d:\Microproject\reclaim\docs\Mitali_Mali_ReClaim_Complete_Project_Report.docx"
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title = doc.add_paragraph()
    r = title.add_run("ReClaim — Smart Lost & Found Management System")
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    r = sub.add_run("Microservices Architecture — Command Execution & Output Report")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(3, 105, 161)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    student = doc.add_paragraph()
    r = student.add_run("Submitted by: Mitali Mali | Enrollment No.: 24162581023\nB.Tech Computer Science & Engineering (Semester 5) | Academic Year 2026")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235)
    student.alignment = WD_ALIGN_PARAGRAPH.CENTER
    student.paragraph_format.space_after = Pt(16)

    # Table
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Student Name", "Mitali Mali"),
        ("Enrollment Number", "24162581023"),
        ("Project Title", "ReClaim — Smart Lost & Found Management System"),
        ("Technical Architecture", "Decoupled Microservices, Docker, Kubernetes, Helm, Postman")
    ]
    for idx, (k, v) in enumerate(meta):
        row = table.rows[idx]
        row.cells[0].text = k
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        set_cell_background(row.cells[0], "F1F5F9")
        row.cells[1].text = v
        set_cell_background(row.cells[1], "FFFFFF")

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph()
    r = p.add_run("Note: Detailed visual terminal screenshots and step-by-step command outputs for all 24 phases are rendered in the companion PDF: Mitali_Mali_ReClaim_Complete_Project_Report.pdf")
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(docx_path)
    print("Generated DOCX:", docx_path)

if __name__ == "__main__":
    create_mitali_docx()
