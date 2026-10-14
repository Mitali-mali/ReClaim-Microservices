import os
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_hindi_presentation_docs():
    html_path = r"d:\Microproject\reclaim\docs\ReClaim_Faculty_Presentation_Guide_Hindi.html"
    pdf_path = r"d:\Microproject\reclaim\docs\ReClaim_Faculty_Presentation_Guide_Hindi.pdf"
    docx_path = r"d:\Microproject\reclaim\docs\ReClaim_Faculty_Presentation_Guide_Hindi.docx"

    # 1. GENERATE HTML FOR PDF
    html_content = """<!DOCTYPE html>
<html lang="hi">
<head>
<meta charset="UTF-8">
<title>ReClaim — Faculty Presentation & Viva Guide (Hindi)</title>
<style>
  @page {
    size: A4;
    margin: 14mm 14mm 14mm 14mm;
    @bottom-right {
      content: counter(page);
    }
  }
  body {
    font-family: 'Segoe UI', Arial, sans-serif;
    color: #1e293b;
    line-height: 1.6;
    background: #ffffff;
    font-size: 10pt;
  }
  .header-box {
    text-align: center;
    border-bottom: 3px solid #2563eb;
    padding-bottom: 14px;
    margin-bottom: 20px;
  }
  h1 {
    color: #1e3a8a;
    font-size: 20pt;
    margin: 0 0 6px 0;
    font-weight: 800;
  }
  .subtitle {
    color: #0369a1;
    font-size: 12pt;
    font-weight: 700;
    margin: 0 0 6px 0;
  }
  .meta {
    color: #64748b;
    font-size: 9pt;
  }
  .section-title {
    color: #1e40af;
    font-size: 13pt;
    font-weight: 800;
    border-left: 5px solid #2563eb;
    padding-left: 10px;
    margin-top: 22px;
    margin-bottom: 10px;
    background: #eff6ff;
    padding-top: 4px;
    padding-bottom: 4px;
  }
  .dialogue-box {
    background: #f8fafc;
    border: 1px solid #cbd5e1;
    border-left: 4px solid #f59e0b;
    padding: 10px 14px;
    border-radius: 4px;
    margin-bottom: 12px;
    font-style: normal;
  }
  .dialogue-speaker {
    font-weight: bold;
    color: #b45309;
    font-size: 9.5pt;
    margin-bottom: 4px;
  }
  .card {
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    margin-bottom: 14px;
    background: #ffffff;
    page-break-inside: avoid;
  }
  .card-header {
    background: #f1f5f9;
    padding: 8px 12px;
    font-weight: bold;
    color: #0f172a;
    border-bottom: 1px solid #e2e8f0;
    font-size: 10.5pt;
  }
  .card-body {
    padding: 10px 12px;
  }
  pre {
    margin: 4px 0 8px 0;
    padding: 8px 10px;
    background: #0f172a;
    color: #38bdf8;
    border-radius: 4px;
    font-family: Consolas, monospace;
    font-size: 9pt;
    white-space: pre-wrap;
    word-break: break-all;
  }
  .why-box {
    background: #ecfdf5;
    border-left: 3px solid #10b981;
    color: #065f46;
    padding: 6px 10px;
    font-size: 9pt;
    margin-bottom: 6px;
  }
  .speak-box {
    background: #fefce8;
    border-left: 3px solid #eab308;
    color: #713f12;
    padding: 6px 10px;
    font-size: 9pt;
  }
  table {
    width: 100%;
    border-collapse: collapse;
    margin-top: 10px;
    margin-bottom: 15px;
    font-size: 9pt;
  }
  table, th, td {
    border: 1px solid #cbd5e1;
  }
  th {
    background: #f1f5f9;
    color: #1e293b;
    font-weight: 700;
    padding: 6px 8px;
    text-align: left;
  }
  td {
    padding: 6px 8px;
    vertical-align: top;
  }
  .highlight {
    font-weight: bold;
    color: #2563eb;
  }
</style>
</head>
<body>

<div class="header-box">
  <h1>ReClaim — Faculty Presentation & Viva Guide</h1>
  <div class="subtitle">Faculty Ke Samne Project Ko Step-by-Step Hindi Me Kaise Explain Karein</div>
  <div class="meta">
    <strong>B.Tech CSE Semester 5</strong> | Microservices Architecture Major Project | <strong>Role:</strong> Solo Presenter
  </div>
</div>

<div class="section-title">1. Faculty Ke Samne Project Ka Opening Pitch (Pehele 2 Minute Me Kya Bolein)</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">🎤 AAPKO AISA BOLNA HAI (Opening Speech):</div>
  "Good morning / Good afternoon Sir/Ma'am. <br>
  Mera naam <strong>[Aapka Naam]</strong> hai, aur mera Semester 5 Microservices Major Project hai <strong>'ReClaim — Smart Lost & Found Management System'</strong>.<br><br>
  College campuses aur public places me aksar students ke important items jaise laptops, ID cards, keys, aur documents kho jate hain. Existing systems monolithic hote hain aur unme manual search karna padta hai.<br><br>
  Maine is problem ko solve karne ke liye ek <strong>Cloud-Native Microservices Architecture</strong> design kiya hai jo hamare syllabus ke charo units ko completely demonstrate karta hai:
  <ul>
    <li><strong>Unit 1 (SOA & REST):</strong> Service-Oriented Architecture jisme har service loosely coupled hai aur pure JSON REST APIs use karti hai.</li>
    <li><strong>Unit 2 (Node.js & Express):</strong> Saari backend services Node.js aur Express me independent modular pattern pe bani hain.</li>
    <li><strong>Unit 3 (Docker & Compose):</strong> Har microservice ka apna lightweight Alpine Dockerfile hai aur master <code>docker-compose.yml</code> se multi-container orchestration hoti hai.</li>
    <li><strong>Unit 4 (Kubernetes & Helm):</strong> Production deployments ke liye Kubernetes manifests (Deployments, ClusterIP/NodePort Services, ConfigMaps, Secrets, Liveness/Readiness probes) aur 1-click deployment ke liye <strong>Helm Chart</strong> create kiya hai."</li>
  </ul>
</div>

<div class="section-title">2. Architecture & Service Breakdown (Screen Par Diagram Dikhate Waqt)</div>

<table>
  <tr>
    <th style="width: 20%;">Service Name</th>
    <th style="width: 12%;">Port</th>
    <th style="width: 28%;">Kya Kaam Karti Hai? (Responsibility)</th>
    <th style="width: 40%;">Faculty Ko Kya Batana Hai? (Technical Logic)</th>
  </tr>
  <tr>
    <td><strong>API Gateway</strong></td>
    <td><code>8080</code></td>
    <td>Single Entry Point, Reverse Proxy, Security</td>
    <td>"Frontend ko har service ka port yaad nahi rakhna padta. Gateway sari incoming requests ko target microservice pe proxy karta hai aur Helmet se security headers lagata hai."</td>
  </tr>
  <tr>
    <td><strong>Auth Service</strong></td>
    <td><code>5001</code></td>
    <td>User Registration, Login, JWT Generation</td>
    <td>"Stateless authentication ke liye JWT (JSON Web Tokens) aur password safety ke liye bcrypt hashing use hoti hai."</td>
  </tr>
  <tr>
    <td><strong>Item Service</strong></td>
    <td><code>5002</code></td>
    <td>Lost & Found Items CRUD & Filtering</td>
    <td>"Users lost ya found items report karte hain. Category, location aur date wise filter karne ke liye optimized queries hain."</td>
  </tr>
  <tr>
    <td><strong>Matching Service</strong></td>
    <td><code>5003</code></td>
    <td>Rule-Based Smart Scoring Engine</td>
    <td>"Bina kisi paid AI ke, pure deterministic algorithm se Category (40%), Location (20%), Color (15%), Brand (15%), aur Date proximity (10%) ka match score nikalta hai."</td>
  </tr>
  <tr>
    <td><strong>Claim Service</strong></td>
    <td><code>5004</code></td>
    <td>Ownership Verification & Workflow</td>
    <td>"Claimant proof deta hai. Reporter/Admin approve ya reject karta hai. Ye service Item service aur Notification service se HTTP se communicate karti hai."</td>
  </tr>
  <tr>
    <td><strong>Notification Service</strong></td>
    <td><code>5005</code></td>
    <td>In-App Alerts Dispatcher</td>
    <td>"Asynchronous alerts store karta hai jab koi match milta hai ya claim status change hota hai."</td>
  </tr>
  <tr>
    <td><strong>MongoDB</strong></td>
    <td><code>27017</code></td>
    <td>Database-per-Service</td>
    <td>"Microservices design pattern ke mutabik har service ka apna isolated logical database hai (e.g. <code>reclaim_auth</code>, <code>reclaim_item</code>), koi direct shared table nahi hai."</td>
  </tr>
</table>

<div class="section-title">3. Live Demo Flow: Kaunsi Command Kab Chalayein, Kyu Chalayein aur Kya Bolein</div>

<div class="card">
  <div class="card-header">STEP 1: Containers Ko Start Karna (Docker Compose)</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu chalana hai:</strong> Pure project ke 8 containers (MongoDB, 5 Microservices, API Gateway, React Frontend) ko ek single command me background me build aur start karne ke liye.</div>
    <pre>cd d:\Microproject\reclaim
docker compose up --build -d</pre>
    <div class="speak-box"><strong>Faculty ko kya bolna hai:</strong> "Sir/Ma'am, main Docker Compose se pure application stack ko instantiate kar rahi hu. Docker automatically ek internal bridge network banata hai jisse sare containers ek dusre ko service name (DNS) se call kar lete hain."</div>
  </div>
</div>

<div class="card">
  <div class="card-header">STEP 2: Containers Ka Health Status Dikhana</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu chalana hai:</strong> Examiner ko proof dikhane ke liye ki sari services active hain, MongoDB healthy state me hai, aur ports sahi bind hue hain.</div>
    <pre>docker compose ps</pre>
    <div class="speak-box"><strong>Faculty ko kya bolna hai:</strong> "Aap dekh sakte hain ki sare 8 containers (reclaim_mongodb, auth, item, matching, claim, notification, api_gateway, frontend) Up aur Healthy hain. Port 3000 frontend ke liye aur 8080 API Gateway ke liye mapped hai."</div>
  </div>
</div>

<div class="card">
  <div class="card-header">STEP 3: Browser Me Live End-to-End Application Demo</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu dikhana hai:</strong> Complete student workflow demonstrate karne ke liye (Registration -> Report -> Smart Match -> Claim -> Notification).</div>
    <pre>Browser me open karein: http://localhost:3000</pre>
    <div class="speak-box"><strong>Faculty ko explain karein:</strong>
      <ol>
        <li><strong>Register/Login:</strong> "Pehle hum user register karenge. Auth service password ko hash karke JWT token return karegi jo browser ke <code>localStorage</code> me save ho jata hai."</li>
        <li><strong>Report Lost Item:</strong> "Hum ek Lost item report karte hain: 'Dell Laptop', Category: 'Electronics', Color: 'Silver', Location: 'Central Library'."</li>
        <li><strong>Report Found Item:</strong> "Dusre student ke roop me Found item report karte hain: 'Dell Inspiron', Category: 'Electronics', Location: 'Library'."</li>
        <li><strong>Smart Matching Engine:</strong> "Item Details page par Matching Service automatically run hoti hai aur 85%+ Match score dikhati hai, sath hi match ke reasons bhi display karti hai."</li>
        <li><strong>Claim & Notification:</strong> "User 'Claim this Item' par click karke proof submit karta hai. Claim Service turant Notification Service ko alert bhejti hai aur item status 'RESOLVED' ho jata hai."</li>
      </ol>
    </div>
  </div>
</div>

<div class="card">
  <div class="card-header">STEP 4: Health Check & Self-Healing Probes Dikhana</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu chalana hai:</strong> Unit 4 (Kubernetes Probes) ka proof dikhane ke liye.</div>
    <pre>curl http://localhost:8080/health
curl http://localhost:5002/health</pre>
    <div class="speak-box"><strong>Faculty ko kya bolna hai:</strong> "Sir/Ma'am, hamari har microservice ek dynamic <code>/health</code> endpoint expose karti hai jo live database connection check karti hai. Agar DB connect hai to HTTP 200 (UP) return hota hai, aur agar DB down ho to HTTP 503 (DEGRADED) return hota hai. Kubernetes isi endpoint ko liveness aur readiness probe ke liye use karta hai."</div>
  </div>
</div>

<div class="card">
  <div class="card-header">STEP 5: Kubernetes (Minikube) Manifests Dikhana</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu chalana hai:</strong> Faculty ko dikhana ki humne actual K8s architecture implement kiya hai (Pods, Deployments, ClusterIP, NodePort, ConfigMap, Secrets).</div>
    <pre>kubectl get pods -n reclaim
kubectl get services -n reclaim
kubectl get configmaps,secrets -n reclaim</pre>
    <div class="speak-box"><strong>Faculty ko kya bolna hai:</strong> "Humne <code>reclaim</code> namespace me 19 Kubernetes manifests deploy kiye hain. Internal services ClusterIP par hain jo bahar se accessible nahi hain, jabki API Gateway aur Frontend NodePort par hain taaki external traffic handle ho sake. ConfigMaps aur Secrets ke jariye humne environment variables decouple kiye hain."</div>
  </div>
</div>

<div class="card">
  <div class="card-header">STEP 6: Helm Packaging Dikhana</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu chalana hai:</strong> Syllabus ka Unit 4 advance topic (Helm Chart Package Manager) prove karne ke liye.</div>
    <pre>cd d:\Microproject\reclaim
helm lint ./helm/reclaim
helm list -n reclaim</pre>
    <div class="speak-box"><strong>Faculty ko kya bolna hai:</strong> "Sir/Ma'am, pure 19 K8s manifests ko manually apply karne ke bajaye humne ek standardized <strong>Helm Chart</strong> banaya hai. Agar hume 10 instances scale karne hain ya environment change karna hai, to sirf <code>values.yaml</code> update karke 1 command se upgrade aur rollback kar sakte hain."</div>
  </div>
</div>

<div class="card">
  <div class="card-header">STEP 7: Postman Collection Test Suite Dikhana</div>
  <div class="card-body">
    <div class="why-box"><strong>Kyu dikhana hai:</strong> Testing standards prove karne ke liye.</div>
    <pre>Postman me import karein:
postman/ReClaim_API_Collection.json
postman/ReClaim_Environment.json</pre>
    <div class="speak-box"><strong>Faculty ko kya bolna hai:</strong> "Humne sare 15 endpoints ke liye Postman test suite banaya hai. Isme JavaScript test scripts hain jo login karte hi token ko extract karke <code>{{token}}</code> environment variable me set kar dete hain, jisse automated regression testing hoti hai."</div>
  </div>
</div>

<div class="section-title">4. Faculty Ke Top 10 Cross-Questions Aur Unke Solid Hindi Answers</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q1. Monolith ke bajaye Microservices kyu choose kiya?</div>
  <strong>Aapka Answer:</strong> "Sir/Ma'am, Monolith me agar search module ya matching calculation fail ho jaye to pura login aur portal crash ho jata tha. Microservices me <strong>Fault Isolation</strong> milta hai — agar Matching Service heavy compute load me hai tab bhi Auth aur Item Service smoothly chalti rahti hain. Sath hi har service ko independently scale kiya ja sakta hai."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q2. Microservices aapas me baat kaise karti hain? (Inter-service communication)</div>
  <strong>Aapka Answer:</strong> "Humne synchronous REST APIs over HTTP use kiya hai. For example, Matching Service ko jab items compare karne hote hain to wo <code>axios.get('http://item-service:5002/api/items')</code> call karti hai. Aise hi Claim Service notification trigger karne ke liye Notification Service ko POST request bhejti hai."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q3. Kya sabhi services ek hi MongoDB database share karti hain?</div>
  <strong>Aapka Answer:</strong> "Nahi Sir/Ma'am! Humne strict <strong>Database-per-Service</strong> pattern follow kiya hai. Auth service ke paas <code>reclaim_auth</code> hai, Item service ke paas <code>reclaim_item</code> hai, aur Matching service ke paas <code>reclaim_matching</code> hai. Koi bhi service dusri service ke database tables ko direct touch nahi kar sakti, sirf APIs ke through data milta hai."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q4. Smart Matching me AI use kiya hai kya?</div>
  <strong>Aapka Answer:</strong> "Nahi Sir/Ma'am, humne honesty ke sath project me <strong>Rule-Based Heuristic Algorithm</strong> implement kiya hai. AI model college laptop par heavy ho jata aur paid cloud APIs chahiye hoti. Hamara algorithm 5 factors par mathematically 0 se 100% score calculate karta hai: Category (40%), Location (20%), Color (15%), Brand (15%), aur Date proximity (10%)."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q5. API Gateway kyu lagaya? Direct service ko call kyu nahi kiya?</div>
  <strong>Aapka Answer:</strong> "Agar Gateway na hota to React frontend ko har service ka port (5001, 5002, 5003...) aur IP address yaad rakhna padta, jisse tight coupling ho jati. Gateway single point of entry provide karta hai port 8080 par, CORS issue solve karta hai, aur Helmet security headers apply karta hai."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q6. Virtual Machine (VM) aur Docker Container me kya farak hai?</div>
  <strong>Aapka Answer:</strong> "VM hardware level par virtualize hoti hai jisme har VM ka apna complete Heavy Guest OS (GBs me) hota hai aur Hypervisor chahiye hota hai. Jabki Docker Container OS-level par virtualize hota hai aur host kernel share karta hai. Isiliye hamare Node Alpine containers sirf ~50MB ke hain aur 2 second me start ho jate hain."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q7. Dockerfile me multi-stage build kyu use kiya?</div>
  <strong>Aapka Answer:</strong> "Frontend ke Dockerfile me humne Multi-Stage build use kiya. Stage 1 me Node.js React code ko build karke static HTML/JS banata hai, aur Stage 2 me hum heavy Node runtime ko discard karke sirf lightweight Nginx server me files copy karte hain. Isse final image size 80% chhota ho jata hai."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q8. Kubernetes me Pod, Deployment aur Service me kya antar hai?</div>
  <strong>Aapka Answer:</strong> 
  <ul>
    <li><strong>Pod:</strong> K8s ka smallest deployable unit jo container ko wrap karta hai.</li>
    <li><strong>Deployment:</strong> Declares desired state (jaise kitne replicas chalane hain, rolling update kaise karna hai).</li>
    <li><strong>Service:</strong> Pods ke ephemeral (temporary) IPs ke upar ek stable static IP aur DNS provide karta hai taaki traffic load-balance ho sake.</li>
  </ul>
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q9. Liveness Probe aur Readiness Probe me kya farak hai?</div>
  <strong>Aapka Answer:</strong> "Readiness Probe check karta hai ki container user requests lene ke liye ready hai ya nahi (e.g. MongoDB connect hui ya nahi). Jab tak ready nahi hota, K8s us Pod ko traffic nahi bhejta. Liveness Probe check karta hai ki container chalte chalte deadlock ya freeze to nahi ho gaya; agar freeze ho gaya to K8s us container ko auto-restart kar deta hai."
</div>

<div class="dialogue-box">
  <div class="dialogue-speaker">Q10. Helm ka kya role hai is project me?</div>
  <strong>Aapka Answer:</strong> "Helm Kubernetes ka Package Manager hai (jaise Node.js ke liye npm hota hai). Bina Helm ke hume 19 alag-alag YAML files bar-bar run karni padti. Helm chart se hum pure ReClaim project ko ek command me install, upgrade, ya rollback kar sakte hain."
</div>

<div class="section-title">5. Presentation Ke Liye Quick Checklist (Apne Sath Kya Ready Rakhein)</div>
<ul>
  <li>✅ Docker Desktop open rakhein taaki <code>docker compose ps</code> turant chale.</li>
  <li>✅ Browser me <code>http://localhost:3000</code> ka tab pehle se open rakhein.</li>
  <li>✅ Terminal me project folder <code>d:\Microproject\reclaim</code> open rakhein.</li>
  <li>✅ Postman me <code>ReClaim_API_Collection.json</code> import karke rakhein.</li>
  <li>✅ Confident rahein — aapne 100% working solo microservices project banaya hai!</li>
</ul>

</body>
</html>
"""

    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {html_path}")

    # Convert to PDF via headless Edge
    edge_path = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    if not os.path.exists(edge_path):
        edge_path = r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"
    if not os.path.exists(edge_path):
        edge_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

    cmd = [
        edge_path,
        "--headless",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        html_path
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode == 0:
        print(f"Successfully generated PDF: {pdf_path}")
    else:
        print(f"PDF error: {res.stderr}")

    # 2. GENERATE DOCX FILE
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    r = title.add_run("ReClaim — Faculty Presentation & Viva Guide (Hindi)")
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    r = sub.add_run("Faculty Ke Samne Project Ko Step-by-Step Hindi Me Kaise Explain Karein\nB.Tech CSE Semester 5 Microservices Major Project")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(3, 105, 161)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1
    h = doc.add_paragraph()
    r = h.add_run("1. Faculty Ke Samne Opening Pitch (Pehele 2 Minute Me Kya Bolna Hai):")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 64, 175)

    p = doc.add_paragraph()
    p.add_run("""Good morning / Good afternoon Sir/Ma'am.
Mera naam [Aapka Naam] hai, aur mera Semester 5 Microservices Major Project hai 'ReClaim — Smart Lost & Found Management System'.

College campuses aur public places me aksar students ke important items jaise laptops, ID cards, keys, aur documents kho jate hain. Existing systems monolithic hote hain aur unme manual search karna padta hai.

Maine is problem ko solve karne ke liye ek Cloud-Native Microservices Architecture design kiya hai jo hamare syllabus ke charo units ko completely demonstrate karta hai:
1. Unit 1 (SOA & REST): Service-Oriented Architecture jisme har service loosely coupled hai aur pure JSON REST APIs use karti hai.
2. Unit 2 (Node.js & Express): Saari backend services Node.js aur Express me independent modular pattern pe bani hain.
3. Unit 3 (Docker & Compose): Har microservice ka apna lightweight Alpine Dockerfile hai aur master docker-compose.yml se multi-container orchestration hoti hai.
4. Unit 4 (Kubernetes & Helm): Production deployments ke liye Kubernetes manifests (Deployments, ClusterIP/NodePort Services, ConfigMaps, Secrets, Liveness/Readiness probes) aur 1-click deployment ke liye Helm Chart create kiya hai.""")

    doc.save(docx_path)
    print(f"Successfully generated DOCX: {docx_path}")

if __name__ == "__main__":
    generate_hindi_presentation_docs()
